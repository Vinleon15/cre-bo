"""Самопроверка без обращения к API: python -m tokensave.selftest

Проверяет главное — что экономия не выбрасывает данные. Три ловушки, в
каждой ответ лежит ровно в одном месте, которое обычно теряется при
оптимизации. Если хоть одна не проходит, сжатие использовать нельзя.
"""

from __future__ import annotations

import tempfile
from pathlib import Path
from types import SimpleNamespace

from . import cache, files, guard, history, store


class FakeClient:
    """Подменяет Anthropic: ничего не отправляет, запоминает собранный запрос."""

    def __init__(self, reply: str = "ответ"):
        self.reply = reply
        self.last = {}
        self.messages = SimpleNamespace(
            create=self._create, count_tokens=self._count)

    def _create(self, **kw):
        self.last = kw
        return SimpleNamespace(
            content=[SimpleNamespace(type="text", text=self.reply)],
            usage=SimpleNamespace(input_tokens=100, output_tokens=10,
                                  cache_read_input_tokens=0,
                                  cache_creation_input_tokens=0),
        )

    def _count(self, **kw):
        return SimpleNamespace(input_tokens=999)

    def context(self) -> str:
        blocks = self.last.get("system", [])
        text = "\n".join(b["text"] for b in blocks)
        for m in self.last.get("messages", []):
            content = m["content"]
            text += "\n" + (content if isinstance(content, str)
                            else "\n".join(b.get("text", "") for b in content))
        return text


def _saver(tmp: Path, client, mode="economy"):
    from .client import Saver
    return Saver(mode=mode, session="t", db=tmp / "test.db", client=client)


def trap_hidden_sheet(tmp: Path) -> tuple[bool, str]:
    """Ключевое число лежит на СКРЫТОМ листе Excel."""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active["A1"] = "Выручка"
    wb.active["B1"] = 1000
    hidden = wb.create_sheet("Допущения")
    hidden["A1"] = "Ставка дисконтирования"
    hidden["B1"] = "13,7%"
    hidden.sheet_state = "hidden"
    path = tmp / "model.xlsx"
    wb.save(path)

    client = FakeClient()
    _saver(tmp, client).ask("какая ставка?", files=[str(path)])
    ctx = client.context()
    return "13,7%" in ctx, "скрытый лист Excel"


def trap_speaker_notes(tmp: Path) -> tuple[bool, str]:
    """Ключевое число только в заметках докладчика, на слайде его нет."""
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Итоги квартала"
    slide.notes_slide.notes_text_frame.text = "Фактический EBITDA 47 300 000 руб."
    path = tmp / "deck.pptx"
    prs.save(path)

    client = FakeClient()
    _saver(tmp, client).ask("какая EBITDA?", files=[str(path)])
    return "47 300 000" in client.context(), "заметки докладчика PPTX"


def trap_old_message(tmp: Path) -> tuple[bool, str]:
    """Ключевое условие названо 20 ходов назад и давно вытеснено из окна."""
    conn = store.connect(tmp / "old.db")
    history.append_turn(conn, "s", "user", "Штраф за просрочку — 0,1% в день, не более 10%.")
    for i in range(20):
        history.append_turn(conn, "s", "user", f"вопрос {i}")
        history.append_turn(conn, "s", "assistant", f"ответ {i}")

    _, pins, recent = history.pack(conn, "s", "economy", summarizer=lambda p: "сводка")
    in_recent = any("0,1%" in t["content"] for t in recent)
    return ("0,1%" in pins and not in_recent), "вытесненный ход (пины)"


def check_cache_audit() -> tuple[bool, str]:
    """Инвалидаторы кэша должны детектироваться."""
    bad = cache.audit("Сегодня 2026-09-18 14:03, сессия 3f2a1b8c-1111-2222-3333-444455556666")
    return len(bad) >= 2, "аудит инвалидаторов кэша"


def check_override() -> tuple[bool, str]:
    """На расчётах и договорах экономия должна отключаться принудительно."""
    must = ["рассчитай NPV", "сверь суммы в договоре", "сравни версии приложения",
            "проверь цифры", "какой штраф за просрочку"]
    must_not = ["привет", "как дела", "перескажи в двух словах"]
    return (all(guard.forces_full_context(q) for q in must)
            and not any(guard.forces_full_context(q) for q in must_not)), "override режима"


def main() -> int:
    checks = [check_cache_audit, check_override]
    traps = [trap_hidden_sheet, trap_speaker_notes, trap_old_message]
    failed = 0

    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        for fn in checks:
            ok, name = fn()
            print(f"  {'OK  ' if ok else 'СБОЙ'} {name}")
            failed += not ok
        print("\n  Ловушки — данные, которые обычно теряются при оптимизации:")
        for fn in traps:
            try:
                ok, name = fn(tmp)
            except Exception as exc:
                ok, name = False, f"{fn.__name__}: {exc}"
            print(f"  {'OK  ' if ok else 'СБОЙ'} {name}")
            failed += not ok

    print(f"\n{'Всё прошло.' if not failed else f'Провалено проверок: {failed}.'}")
    if failed:
        print("Сжатие использовать нельзя, пока ловушки не проходят.")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
